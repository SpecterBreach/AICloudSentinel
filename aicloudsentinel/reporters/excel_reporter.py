"""
reporters/excel_reporter.py
============================
Generates a professional multi-sheet Excel report from ScanResult.

Sheets:
  1. Executive Summary  — risk dashboard, narrative, scan metadata
  2. All Findings       — full sortable table with OWASP mappings + AI analysis
  3. Priority Actions   — CRITICAL + HIGH only with AI remediation steps
  4. OWASP Heatmap      — finding counts mapped to OWASP LLM + Agentic Top 10
"""

from __future__ import annotations
import logging
from pathlib import Path
from datetime import datetime, timezone

from ..models import ScanResult, Finding, RiskLevel

log = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side, GradientFill
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY    = "0D1B2A"
TEAL    = "1B7A6E"
RED     = "C0392B"
ORANGE  = "E67E22"
YELLOW  = "D4AC0D"
GREEN   = "1D8348"
LGRAY   = "F4F6F7"
WHITE   = "FFFFFF"

RISK_COLORS = {
    "CRITICAL": RED,
    "HIGH":     ORANGE,
    "MEDIUM":   YELLOW,
    "LOW":      GREEN,
    "INFO":     "2980B9",
}

THIN = Border(
    left  = Side(style="thin", color="D5D8DC"),
    right = Side(style="thin", color="D5D8DC"),
    top   = Side(style="thin", color="D5D8DC"),
    bottom= Side(style="thin", color="D5D8DC"),
)

def _hdr(ws, cols: list[str], row: int = 1) -> None:
    for ci, col in enumerate(cols, 1):
        c = ws.cell(row=row, column=ci, value=col)
        c.fill   = PatternFill("solid", fgColor=NAVY)
        c.font   = Font(name="Calibri", bold=True, color=WHITE, size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN
    ws.row_dimensions[row].height = 28

def _autowidth(ws, max_w: int = 55) -> None:
    for col in ws.columns:
        w = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 3, max_w)


class ExcelReporter:
    """
    Writes an AICloudSentinel scan result to a formatted Excel workbook.

    Usage:
        reporter = ExcelReporter()
        path     = reporter.write(scan_result, "report.xlsx")
    """

    def write(self, result: ScanResult, output_path: str | Path = "aicloudsentinel_report.xlsx") -> Path:
        if not _HAS_OPENPYXL:
            raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

        wb   = openpyxl.Workbook()
        path = Path(output_path)

        self._sheet_summary(wb, result)
        self._sheet_findings(wb, result)
        self._sheet_priority(wb, result)
        self._sheet_owasp_heatmap(wb, result)

        wb.save(path)
        log.info("Report saved → %s", path.resolve())
        return path

    # ── Sheet 1: Executive Summary ────────────────────────────────────────────

    def _sheet_summary(self, wb, result: ScanResult) -> None:
        ws = wb.active
        ws.title = "Executive Summary"
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 50

        # Title bar
        ws.merge_cells("B2:C2")
        t = ws["B2"]
        t.value     = "AICloudSentinel — AI Workload Security Report"
        t.font      = Font(name="Calibri", size=18, bold=True, color=NAVY)
        t.alignment = Alignment(vertical="center")
        ws.row_dimensions[2].height = 36

        # Metadata
        meta = [
            ("Cloud Provider",   result.provider.value),
            ("Scan Scope",       result.scope),
            ("Scan Date",        result.started_at[:10]),
            ("Completed",        result.completed_at[:10] if result.completed_at else "—"),
            ("Scan ID",          result.scan_id),
        ]
        for row, (k, v) in enumerate(meta, start=4):
            ws.cell(row=row, column=2, value=k).font = Font(name="Calibri", bold=True, color=NAVY, size=10)
            ws.cell(row=row, column=3, value=v).font = Font(name="Calibri", size=10)

        # Risk KPIs
        ws.merge_cells("B10:C10")
        ws["B10"].value = "Risk Summary"
        ws["B10"].font  = Font(name="Calibri", bold=True, size=12, color=TEAL)

        kpis = [
            ("CRITICAL", result.critical_count, RED),
            ("HIGH",     result.high_count,     ORANGE),
            ("MEDIUM",   result.medium_count,   YELLOW),
            ("LOW",      result.low_count,       GREEN),
            ("TOTAL",    result.total,            NAVY),
        ]
        for ci, (label, count, color) in enumerate(kpis, start=2):
            col = get_column_letter(ci)
            label_cell = ws.cell(row=12, column=ci, value=label)
            count_cell = ws.cell(row=13, column=ci, value=count)
            label_cell.fill      = PatternFill("solid", fgColor=color)
            label_cell.font      = Font(name="Calibri", bold=True, color=WHITE, size=9)
            label_cell.alignment = Alignment(horizontal="center")
            count_cell.font      = Font(name="Calibri", bold=True, size=16, color=color)
            count_cell.alignment = Alignment(horizontal="center")

        ws.row_dimensions[12].height = 20
        ws.row_dimensions[13].height = 28

        # Executive summary text
        ws.merge_cells("B15:C15")
        ws["B15"].value = "Executive Summary"
        ws["B15"].font  = Font(name="Calibri", bold=True, size=12, color=TEAL)

        ws.merge_cells("B16:C22")
        summary_cell = ws["B16"]
        summary_cell.value     = result.executive_summary or "AI analysis not available."
        summary_cell.font      = Font(name="Calibri", size=10)
        summary_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[16].height = 120

        # Risk narrative
        ws.merge_cells("B24:C24")
        ws["B24"].value = "Risk Narrative"
        ws["B24"].font  = Font(name="Calibri", bold=True, size=12, color=TEAL)

        ws.merge_cells("B25:C27")
        narrative_cell = ws["B25"]
        narrative_cell.value     = result.risk_narrative or "—"
        narrative_cell.font      = Font(name="Calibri", size=10)
        narrative_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[25].height = 60

    # ── Sheet 2: All Findings ─────────────────────────────────────────────────

    def _sheet_findings(self, wb, result: ScanResult) -> None:
        ws = wb.create_sheet("All Findings")
        cols = [
            "Finding ID", "Risk Level", "Title", "Provider",
            "Resource Type", "Resource Name", "Region",
            "OWASP LLM", "OWASP Agentic",
            "Description", "Evidence", "Recommendation",
            "AI Attack Scenario", "AI Remediation", "AI Business Impact",
            "Timestamp",
        ]
        _hdr(ws, cols)
        ws.freeze_panes = "A2"

        for ri, f in enumerate(result.sorted_findings(), start=2):
            row_data = [
                f.finding_id,
                f.risk_level.value,
                f.title,
                f.provider.value,
                f.resource_type,
                f.resource_name,
                f.region,
                ", ".join(f.owasp.llm_ids),
                ", ".join(f.owasp.agentic_ids),
                f.description,
                f.evidence,
                f.recommendation,
                f.ai_attack_scenario,
                f.ai_remediation,
                f.ai_explanation,
                f.timestamp[:10],
            ]
            for ci, val in enumerate(row_data, start=1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.border    = THIN
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font      = Font(name="Calibri", size=9)
                if cols[ci - 1] == "Risk Level":
                    cell.fill = PatternFill("solid", fgColor=RISK_COLORS.get(f.risk_level.value, "FFFFFF"))
                    cell.font = Font(name="Calibri", bold=True, color=WHITE, size=9)
                elif ri % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=LGRAY)

        _autowidth(ws)

    # ── Sheet 3: Priority Actions ─────────────────────────────────────────────

    def _sheet_priority(self, wb, result: ScanResult) -> None:
        ws   = wb.create_sheet("Priority Actions")
        cols = [
            "Risk Level", "Title", "Provider", "Resource Type", "Resource Name",
            "OWASP Mapping", "AI Attack Scenario", "AI Remediation Steps",
        ]
        _hdr(ws, cols)
        ws.freeze_panes = "A2"

        priority = [
            f for f in result.sorted_findings()
            if f.risk_level in (RiskLevel.CRITICAL, RiskLevel.HIGH)
        ]

        for ri, f in enumerate(priority, start=2):
            row = [
                f.risk_level.value,
                f.title,
                f.provider.value,
                f.resource_type,
                f.resource_name,
                f.owasp.describe(),
                f.ai_attack_scenario or f.description,
                f.ai_remediation or f.recommendation,
            ]
            for ci, val in enumerate(row, start=1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.border    = THIN
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font      = Font(name="Calibri", size=9)
                if cols[ci - 1] == "Risk Level":
                    cell.fill = PatternFill("solid", fgColor=RISK_COLORS.get(f.risk_level.value, "FFFFFF"))
                    cell.font = Font(name="Calibri", bold=True, color=WHITE, size=9)
                elif ri % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=LGRAY)

        _autowidth(ws)

    # ── Sheet 4: OWASP Heatmap ────────────────────────────────────────────────

    def _sheet_owasp_heatmap(self, wb, result: ScanResult) -> None:
        from ..models import OWASP_LLM_TOP10, OWASP_AGENTIC_TOP10

        ws = wb.create_sheet("OWASP Heatmap")
        ws.sheet_view.showGridLines = False

        ws["B2"].value = "OWASP LLM Top 10 — Finding Counts"
        ws["B2"].font  = Font(name="Calibri", bold=True, size=12, color=NAVY)

        llm_counts: dict[str, int] = {}
        agt_counts: dict[str, int] = {}
        for f in result.findings:
            for lid in f.owasp.llm_ids:
                llm_counts[lid] = llm_counts.get(lid, 0) + 1
            for aid in f.owasp.agentic_ids:
                agt_counts[aid] = agt_counts.get(aid, 0) + 1

        _hdr(ws, ["ID", "Category", "Count"], row=3)
        for ri, (oid, name) in enumerate(OWASP_LLM_TOP10.items(), start=4):
            cnt = llm_counts.get(oid, 0)
            ws.cell(row=ri, column=2, value=oid).font   = Font(name="Calibri", bold=True, size=9, color=TEAL)
            ws.cell(row=ri, column=3, value=name).font  = Font(name="Calibri", size=9)
            c = ws.cell(row=ri, column=4, value=cnt)
            c.font      = Font(name="Calibri", bold=True, size=9,
                               color=WHITE if cnt > 0 else NAVY)
            c.fill      = PatternFill("solid", fgColor=ORANGE if cnt > 0 else LGRAY)
            c.alignment = Alignment(horizontal="center")

        # Agentic Top 10
        ws.cell(row=16, column=2).value = "OWASP Agentic AI Top 10 — Finding Counts"
        ws.cell(row=16, column=2).font  = Font(name="Calibri", bold=True, size=12, color=NAVY)
        _hdr(ws, ["ID", "Category", "Count"], row=17)
        for ri, (oid, name) in enumerate(OWASP_AGENTIC_TOP10.items(), start=18):
            cnt = agt_counts.get(oid, 0)
            ws.cell(row=ri, column=2, value=oid).font   = Font(name="Calibri", bold=True, size=9, color=RED)
            ws.cell(row=ri, column=3, value=name).font  = Font(name="Calibri", size=9)
            c = ws.cell(row=ri, column=4, value=cnt)
            c.font      = Font(name="Calibri", bold=True, size=9,
                               color=WHITE if cnt > 0 else NAVY)
            c.fill      = PatternFill("solid", fgColor=RED if cnt > 0 else LGRAY)
            c.alignment = Alignment(horizontal="center")

        ws.column_dimensions["B"].width = 10
        ws.column_dimensions["C"].width = 45
        ws.column_dimensions["D"].width = 12
